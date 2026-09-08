from __future__ import annotations

import re
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from ..csv_utils import decode_bytes, normalize_header, parse_datetime, rows_from_csv
from ..isin import country_from_isin, infer_from_xtb_symbol
from ..models import ParsedFile, Transaction, TxType
from ..money import D, money

CLOSED_HINTS = ("closed position", "closed positions", "historia pozycji", "pozycje zamknięte")
CASH_HINTS = ("cash operation", "cash operations", "operacje gotówkowe", "historia operacji")


def looks_like_xtb(headers: list[str], text: str = "") -> bool:
    blob = " ".join(headers) + " " + text[:800].lower()
    header_set = set(headers)
    if {"type", "amount", "symbol"} <= header_set and ("time" in header_set or "comment" in header_set):
        return True
    return any(
        token in blob
        for token in (
            "open time",
            "close time",
            "open price",
            "gross p/l",
            "closed position",
            "cash operation",
            "purchase value",
            "divident",
            "withholding tax",
            "free-funds interest",
        )
    ) or ("symbol" in header_set and "volume" in header_set)


def parse_xtb(filename: str, data: bytes, account_currency: str = "PLN") -> ParsedFile:
    name = filename.lower()
    if name.endswith((".xlsx", ".xls")):
        return _parse_xlsx(filename, data, account_currency)
    return _parse_csv(filename, decode_bytes(data), account_currency)


def _parse_xlsx(filename: str, data: bytes, account_currency: str) -> ParsedFile:
    workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    transactions: list[Transaction] = []
    skipped: list[str] = []
    notes: list[str] = []
    detected = False
    for sheet in workbook.worksheets:
        rows = _sheet_rows(sheet)
        kind = _sheet_kind(sheet.title, rows)
        if kind == "closed":
            detected = True
            parsed, skip = _closed_from_rows(rows, account_currency, f"{filename}:{sheet.title}")
            transactions.extend(parsed)
            skipped.extend(skip)
        elif kind == "cash":
            detected = True
            parsed, skip = _cash_from_rows(rows, account_currency, f"{filename}:{sheet.title}")
            transactions.extend(parsed)
            skipped.extend(skip)
    if not detected:
        for sheet in workbook.worksheets:
            rows = _sheet_rows(sheet)
            parsed, skip = _closed_from_rows(rows, account_currency, f"{filename}:{sheet.title}")
            if parsed:
                transactions.extend(parsed)
                skipped.extend(skip)
                detected = True
    if not transactions:
        notes.append("XTB workbook had no closed-position or cash-operation rows we could read.")
    return ParsedFile(filename=filename, broker="xtb", transactions=transactions, skipped=skipped, notes=notes)


def _parse_csv(filename: str, text: str, account_currency: str) -> ParsedFile:
    headers, views = rows_from_csv(text)
    header_blob = " ".join(headers).lower() + text[:400].lower()
    raw_rows = [_row_dict(view, headers) for view in views]
    if any(h in header_blob for h in CASH_HINTS) or {"type", "amount"} <= set(normalize_header(h) for h in headers) and "volume" not in header_blob:
        parsed, skipped = _cash_from_dicts(raw_rows, account_currency, filename)
        return ParsedFile(filename=filename, broker="xtb", transactions=parsed, skipped=skipped)
    parsed, skipped = _closed_from_dicts(raw_rows, account_currency, filename)
    return ParsedFile(filename=filename, broker="xtb", transactions=parsed, skipped=skipped)


def _sheet_rows(sheet) -> list[list[object]]:
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _sheet_kind(title: str, rows: list[list[object]]) -> str | None:
    blob = title.lower()
    if any(h in blob for h in CLOSED_HINTS):
        return "closed"
    if any(h in blob for h in CASH_HINTS):
        return "cash"
    preview = " ".join(_stringify(cell) for row in rows[:8] for cell in row).lower()
    if any(h in preview for h in CLOSED_HINTS):
        return "closed"
    if any(h in preview for h in CASH_HINTS):
        return "cash"
    if "open time" in preview and "symbol" in preview:
        return "closed"
    if "amount" in preview and "type" in preview and "comment" in preview:
        return "cash"
    return None


def _closed_from_rows(rows: list[list[object]], account_currency: str, source: str) -> tuple[list[Transaction], list[str]]:
    header_idx, mapping = _find_header(rows, ("symbol", "open time", "volume", "type", "open price"))
    if header_idx is None:
        return [], []
    dicts = _rows_to_dicts(rows[header_idx + 1 :], mapping)
    return _closed_from_dicts(dicts, account_currency, source)


def _cash_from_rows(rows: list[list[object]], account_currency: str, source: str) -> tuple[list[Transaction], list[str]]:
    header_idx, mapping = _find_header(rows, ("type", "time", "amount"))
    if header_idx is None:
        return [], []
    dicts = _rows_to_dicts(rows[header_idx + 1 :], mapping)
    return _cash_from_dicts(dicts, account_currency, source)


def _find_header(rows: list[list[object]], required: tuple[str, ...]) -> tuple[int | None, dict[str, int]]:
    needed = set(required)
    for idx, row in enumerate(rows[:40]):
        mapping: dict[str, int] = {}
        for col, cell in enumerate(row):
            key = normalize_header(_stringify(cell))
            if key:
                mapping[key] = col
        if needed <= set(mapping):
            return idx, mapping
        if {"symbol", "volume", "open time"} <= set(mapping):
            return idx, mapping
        if {"type", "amount"} <= set(mapping) and ("time" in mapping or "comment" in mapping):
            return idx, mapping
    return None, {}


def _rows_to_dicts(rows: list[list[object]], mapping: dict[str, int]) -> list[dict[str, str]]:
    result = []
    for row in rows:
        if not any(_stringify(cell) for cell in row):
            continue
        item = {key: _stringify(row[col]) if col < len(row) else "" for key, col in mapping.items()}
        if any(item.values()):
            result.append(item)
    return result


def _closed_from_dicts(rows: list[dict[str, str]], account_currency: str, source: str) -> tuple[list[Transaction], list[str]]:
    transactions: list[Transaction] = []
    skipped: list[str] = []
    for index, row in enumerate(rows, start=1):
        symbol = _pick(row, "symbol")
        if not symbol or symbol.lower() in {"total", "symbol"}:
            continue
        try:
            open_ts = _as_datetime(_pick(row, "open time", "time"))
            close_raw = _pick(row, "close time")
            close_ts = _as_datetime(close_raw) if close_raw else open_ts + timedelta(seconds=1)
        except ValueError as exc:
            skipped.append(f"{source} row {index}: {exc}")
            continue
        volume = abs(D(_pick(row, "volume", "quantity", "qty")))
        if volume <= 0:
            continue
        open_price = abs(D(_pick(row, "open price")))
        close_price = abs(D(_pick(row, "close price", "market price")))
        commission = abs(D(_pick(row, "commission")))
        swap = D(_pick(row, "swap"))
        position = _pick(row, "position", "id", "order") or f"xtb-{index}"
        inferred_ccy, inferred_country = infer_from_xtb_symbol(symbol)
        currency = (_pick(row, "currency") or inferred_ccy or account_currency).upper()
        side = _pick(row, "type", "side").upper()
        fee_open = money(commission / 2) if commission else D(0)
        fee_close = money(commission - fee_open + abs(swap)) if (commission or swap) else D(0)
        transactions.append(
            Transaction(
                id=f"{position}-open",
                broker="xtb",
                timestamp=open_ts,
                type=TxType.BUY,
                symbol=symbol,
                quantity=volume,
                price=open_price,
                currency=currency,
                fee=fee_open,
                fee_currency=account_currency,
                country=inferred_country or country_from_isin(_pick(row, "isin") or None),
                notes=_pick(row, "comment"),
                raw_action=f"open {side or 'BUY'}",
            )
        )
        transactions.append(
            Transaction(
                id=f"{position}-close",
                broker="xtb",
                timestamp=close_ts,
                type=TxType.SELL,
                symbol=symbol,
                quantity=volume,
                price=close_price,
                currency=currency,
                fee=fee_close,
                fee_currency=account_currency,
                country=inferred_country,
                notes=_pick(row, "comment"),
                raw_action=f"close {side or 'BUY'}",
            )
        )
    return transactions, skipped


def _cash_from_dicts(rows: list[dict[str, str]], account_currency: str, source: str) -> tuple[list[Transaction], list[str]]:
    transactions: list[Transaction] = []
    skipped: list[str] = []
    pending_div: dict[str, Transaction] = {}
    for index, row in enumerate(rows, start=1):
        op = _pick(row, "type").strip()
        if not op:
            continue
        try:
            ts = _as_datetime(_pick(row, "time"))
        except ValueError as exc:
            skipped.append(f"{source} cash row {index}: {exc}")
            continue
        amount = D(_pick(row, "amount"))
        symbol = _pick(row, "symbol")
        comment = _pick(row, "comment")
        op_id = _pick(row, "id") or f"xtb-cash-{index}"
        key = op.lower()
        inferred_ccy, inferred_country = infer_from_xtb_symbol(symbol) if symbol else (account_currency, None)
        if "dividen" in key:
            tx = Transaction(
                id=op_id,
                broker="xtb",
                timestamp=ts,
                type=TxType.DIVIDEND,
                symbol=symbol,
                quantity=D(1),
                price=abs(amount),
                currency=account_currency,
                country=inferred_country,
                notes=comment,
                raw_action=op,
            )
            transactions.append(tx)
            pending_div[_div_key(symbol, ts)] = tx
        elif "withholding" in key or "podatek u źródła" in key or "withholding tax" in key:
            match = pending_div.get(_div_key(symbol, ts))
            if match:
                match.withholding_tax = abs(amount)
                match.withholding_currency = account_currency
            else:
                transactions.append(
                    Transaction(
                        id=op_id,
                        broker="xtb",
                        timestamp=ts,
                        type=TxType.DIVIDEND,
                        symbol=symbol,
                        quantity=D(1),
                        price=D(0),
                        currency=account_currency,
                        withholding_tax=abs(amount),
                        withholding_currency=account_currency,
                        country=inferred_country,
                        notes=comment,
                        raw_action=op,
                    )
                )
        elif "interest" in key and "tax" not in key:
            transactions.append(
                Transaction(
                    id=op_id,
                    broker="xtb",
                    timestamp=ts,
                    type=TxType.INTEREST,
                    symbol=symbol,
                    price=abs(amount),
                    currency=account_currency,
                    notes=comment,
                    raw_action=op,
                )
            )
        elif "interest tax" in key:
            continue
        elif key.startswith("deposit") or key.startswith("withdrawal"):
            continue
        else:
            skipped.append(f"{source}: ignored cash type '{op}'")
    return transactions, skipped


def _div_key(symbol: str, ts: datetime) -> str:
    return f"{symbol}|{ts.date().isoformat()}"


def _pick(row: dict[str, str], *names: str) -> str:
    normalized = {normalize_header(k): v for k, v in row.items()}
    for name in names:
        key = normalize_header(name)
        if key in normalized and normalized[key] != "":
            return str(normalized[key]).strip()
    return ""


def _as_datetime(value: object) -> datetime:
    if isinstance(value, datetime):
        return value
    text = _stringify(value)
    if not text:
        raise ValueError("missing datetime")
    if re.fullmatch(r"\d+(\.\d+)?", text):
        serial = float(text)
        return datetime(1899, 12, 30) + timedelta(days=serial)
    return parse_datetime(text)


def _stringify(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value).strip()


def _row_dict(view, headers: list[str]) -> dict[str, str]:
    return {normalize_header(h): view.get(h) for h in headers}
